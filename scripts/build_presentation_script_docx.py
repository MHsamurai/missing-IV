from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "manuscript" / "vector_missing_iv_presentation_script_ja.docx"
LATIN_FONT = "Arial"
JAPANESE_FONT = "BIZ UDPGothic"


SLIDES = [
    ("Slide 1  Title", [
        "本報告の題目は、欠測IVと潜在変数モデリングによる非無作為欠測の識別です。多次元の測定項目がmissing not at random、すなわちMNARで欠測する状況で、潜在分布と測定核をどのように識別するかを議論します。",
        "中心となる考え方は二段階です。第一段階でMissing IVとsupported blockを用いて観測変数のblock lawを回復し、第二段階でanchor itemsとlatent shifterを用いて、その法則を有限潜在クラスへ一意に分解します。",
    ]),
    ("Slide 2  Missing-data mechanisms and standard approaches", [
        "まず欠測機構を整理します。Dをresponse pattern、Y obsとY misを観測部分と欠測部分とします。MCARではDはYに依存せず、MARでは観測値だけに依存します。MNARではDが欠測値そのものにも依存するため、通常の補完やrespondent-only analysisではsample-selection biasが生じ得ます。",
        "本報告では、欠測のためのinstrumental variableを最初にMissing IVまたはshadow variableと呼び、以後はMissing IVに統一します。Missing IVによってfull-data lawを回復した後に、潜在構造を分解することが課題です。",
    ]),
    ("Slide 3  Usual IV and Missing IV", [
        "通常のIVとMissing IVは役割が異なります。通常のIVはoutcome equationから除外され、treatmentや内生変数を動かします。これに対してMissing IVのZは、欠測指標Rへの直接効果を持たず、欠測し得るYを予測します。",
        "本稿の基本条件は、Yと共変量Uを条件づけるとRとZが独立になることです。Zのrelevanceと、このexclusion、さらにcompletenessを組み合わせて、観測確率を一意に回復します。",
        "パイは、Yの値ごとの観測確率です。その逆数を使って観測された標本を重み付けすることで、欠測した人も含めた母集団全体のYの分布を回復します。",
    ]),
    ("Slide 4  A simple multivariate extension", [
        "単純な多変量拡張では、D jを各項目Y jの観測指標とし、全項目が同時に観測されたcomplete-case indicatorをRと置きます。Rが正の確率を持ち、vector Y全体に対するMissing IV条件とcompletenessが成立すれば、full-data lawを識別できます。",
        "ただし、この方法は全項目を同時に観測するglobal complete caseを必要とします。項目数が多い場合にはこの確率が極端に小さくなり得るため、本稿では後でsupported blocksへ置き換えます。",
    ]),
    ("Slide 5  Latent measurement model", [
        "Fを未観測の潜在変数、Y 1からY mをその測定項目、Xを通常共変量、Wを潜在分布を動かすlatent shifterとします。この構造はstructural equation modeling、SEMやitem response theory、IRTの基礎となる測定モデルです。",
        "Q FはFの潜在分布、Q jは項目jの測定核です。局所独立性の下で、観測変数の分布はQ FとQ jの積をFについて周辺化したmixtureになります。完全データで三つのviewが十分なrankを持つ場合、Allman、Matias and Rhodesの結果により、有限class比率と測定核はlabelを除いて識別されます。",
        "今回の識別定理が直接扱うのは有限潜在クラスです。連続因子のSEMやIRTへ拡張するには、測定族やcalibrationなどの追加条件が必要です。",
    ]),
    ("Slide 6  Related identification results", [
        "関連研究を三つに整理します。scalar Missing IVではd’Haultfoeuille、Zhao and Shao、Miaoらがcompletenessとconditional moment identificationを扱います。multivariate MNARではTangら、Sadinle and Reiter、Liら、Ni and Shaoがpattern restrictionやitemwise nonresponse modelを扱います。",
        "latent structure側ではAllman、Matias and Rhodesのtensor uniquenessに加え、Muthénら、Holman and Glas、Lee and Tang、Harel and Schafer、Jungら、Kano and Takai、Kuhaらがlatent variableとnonignorable missingnessを同時に扱い、Xieらはdeep latent MNARを検討しています。latent-MNARのjoint model自体は既知です。",
        "本稿の新規性はlatent-MNAR自体ではありません。parametric selection modelを置かないsupported-block lawの回復と、共通label下でのlatent decompositionを接続し、一つの識別定理として示す点です。",
    ]),
    ("Slide 7  A latent shifter is not a Missing IV", [
        "latent shifter WとMissing IV Z Sは区別します。仮にFまで条件づければDとWが独立であっても、Fを積分するとDの条件付き分布にはP of F given Y, W, Xが残ります。WがFの分布を動かすなら、この分布は一般にWへ依存します。",
        "したがって、Wは周辺化後のMissing IVにはなりません。Z Sはblock lawをMNAR selectionから回復し、Wは回復後のtensorにthird-mode variationを与えます。block lawを回復しただけではQ FとQ Yは分離されないため、この二つの役割を接続する必要があります。",
    ]),
    ("Slide 8  Model and assumptions", [
        "第一段階では、Missing IV、positivity、complete-case completenessを用いて、観測データから各supported-block lawを回復します。この段階ではlatent mixtureの分解は行いません。",
    ]),
    ("Slide 9  Observed data and supported blocks", [
        "D jをY jの観測指標とし、Oを観測変数の集合とします。Oは常時観測されるX、W、Z、Dと、Dによって選択されたYの成分Y Dから構成されます。R Sは集合Sの全項目が観測されたことを示すblock indicatorです。R SがゼロならY Sは欠測であり、観測データには値を持つ変数として現れません。",
        "supported blockとは、対象support上でR Sが1となる条件付き確率pi Sが正のblockです。必要なのはanchor pairとextension pairsがsupportedであることであり、global complete caseではありません。",
    ]),
    ("Slide 10  Single-block data-generating structure", [
        "図は一つのsupported blockの構造です。U SはXとWの組であり、F、Z S、Y Sと関連します。Z SはFを通じてY Sを予測しますが、Y SとU Sを条件づけた後にはR Sへ直接影響しません。そのためZ SからR Sへの直接辺は置きません。",
        "一方、R Sは欠測し得るY Sに依存してよく、MNARを許容します。ここでのDAGは、本稿が用いる条件付き独立性と分布の関係を整理したものです。",
    ]),
    ("Slide 11  Overlapping supported blocks replace global complete cases", [
        "global complete caseの代わりに、重なりを持つsupported pairsを用います。S 0はanchor itemsのY aとY bを含み、各nonanchor item Y jには、Y jと共通anchor Y aからなるextension pairを用意します。",
        "各pairに固有のMissing IV Z Sとbridgeを置き、第一段階でpair lawを回復します。第二段階では共有されたanchorを通じて、block間のmeasurement kernelとclass labelを接続します。",
    ]),
    ("Slide 12  Stage-1 assumptions", [
        "仮定6はblockwise covariate-adjusted Missing IVです。各supported blockについて、Y SとU Sを条件づけるとR SとZ Sが独立であるとします。",
        "仮定7はcomplete-case completenessです。R Sが1である標本において、Z SとU Sで条件づけたhの期待値がゼロなら、h自体がほとんど確実にゼロであるとします。これは観測可能なcomplete-case conditional operatorの単射性です。positivityはcomplete-case lawとfull lawの零集合を対応させます。",
    ]),
    ("Slide 13  New Proposition 1", [
        "新しい命題1は、本文の命題2.1に対応するsupported-block identificationです。真のpi Sの逆数がcomplete-case lawの下で二乗可積分であると仮定します。その上で、observable bridge momentを満たし、逆数が同じL2空間に属する正値関数のうち、pi Sが唯一の解になります。存在は真のpropensityから与えられ、命題の中心はこのbridge class内の一意性です。",
        "pi Sが識別されると、R S divided by pi Sによるinverse probability weightingで任意のblock関数のfull-data expectationを再現できます。したがって各block lawが、parametric selection modelを指定せずに識別されます。",
    ]),
    ("Slide 14  Proof sketch the observable bridge is unique", [
        "証明では仮定の役割を分けます。観測過程、supported blockの定義、Missing IV exclusionから、真のpi Sがbridge momentを満たすことを示します。別の解との差を取ると、complete cases上の条件付き期待値がゼロになります。",
        "complete-case completenessがその差をゼロにし、positivityが等式を対象support全体へ拡張します。最後にinverse weighting identityからblock lawの識別が従います。",
    ]),
    ("Slide 15  Finite latent-class identification", [
        "第二段階では、第一段階で回復したblock lawsを有限潜在クラスへ分解します。anchor itemsとlatent shifterを用いてKruskal型の一意性を得ることが目的です。",
    ]),
    ("Slide 16  The remaining decomposition problem", [
        "第一段階の後には、必要なY S、W、Xのfull pair lawsが識別されています。しかし潜在変数Fは未観測であり、同じpair lawを生成するp f、G、M jが複数存在する可能性があります。",
        "ここでp fはclass比率、Gはclass別のWの分布、M jはclass別の測定核です。Lewbelの用語では、同じobserved-data lawを生成する二つの値はobservationally equivalentです。anchor orderingを課さないとclass permutationを除く同値類が識別され、orderingを課すと正規化されたparameter space上でpoint identificationを得ます。",
    ]),
    ("Slide 17  Identification strategy", [
        "識別戦略は四段階です。まずsupported-block lawsを回復し、次にanchor tensorとextension tensorsを構成します。anchor tensorを一意に分解してp、G、M a、M bを得た後、class weightsを回復し、各extension tensorを線形反転して全M jを得ます。",
        "つまり第二段階は、回復済みblock lawsをanchor decompositionとextension inversionで潜在分布と測定核へ分ける操作です。",
    ]),
    ("Slide 18  Lemma 3.1", [
        "補題3.1はKhatri–Rao積のrank条件です。AとBはゼロ列を持たないとします。column Kruskal rankの和がrプラス1以上なら、B Khatri–Rao Aはfull column rankになります。",
        "この補題により、anchor tensorのKruskal条件を、各extension itemを線形反転するためのrank条件へ変換できます。",
    ]),
    ("Slide 19  Assumption 3.1", [
        "仮定3.1はcovariate-assisted anchor-pair structureです。class数rは既知です。anchor decompositionには、anchor pairのsupportedness、正のclass比率とWの周辺確率、M a、M b、GのKruskal rank条件を用います。これはAllman、Matias and Rhodesのthree-view conditionと同じです。",
        "既知のanchor scoreの厳密順序で共通labelを固定します。extension propagationでは全extension pairsがsupportedであることと、補題3.1からG Khatri–Rao M aがfull column rankになることを用います。",
    ]),
    ("Slide 20  New Theorem 1", [
        "新しい定理1は、本文の定理3.1に対応するanchor-pair identificationです。第一段階の仮定と仮定3.1の下で、観測データ法則からp、G、Wごとの潜在class分布Q F、全項目の測定核Q Yが固定された共通labelの下で識別されます。",
        "要するに、pair lawsを回復し、一つのanchor tensorを分解すれば、latent-class distributionと全measurement kernelsが得られます。その結果、同時観測されたことのない項目集合についてもjoint lawを構成できます。",
    ]),
    ("Slide 21  Proof map for Theorem 1", [
        "証明の流れを一枚で示します。observed-data lawから命題1でanchor tensorとextension tensorsを回復し、Allman型の一意性からM a、M b、p、Gを共通labelの下で得ます。",
        "次に補題3.1でextension inversionのrankを確保します。pとGにはBayes則を適用してWごとのlatent distributionを得た後、各extension tensorを反転して全M jを回復します。詳細な8ステップはAppendixに置いています。",
    ]),
    ("Slide 22  Bridge-weighted composite estimation", [
        "推定も識別と同じ順序です。第一段階はd’HaultfoeuilleのSection 3に従い、各blockでconditional moment equationを解いてpair lawをIPWで回復します。第二段階はZhao and Shaoのplug-inの順序に従い、共通のlatent mixtureをoverlapping pair lawsへ当てはめます。",
        "提案量はfull likelihoodではなくbridge-weighted composite M-estimatorです。一致性と漸近正規性の定理は有限次元bridgeを対象とし、sandwich varianceにはfirst-stage bridgeの推定誤差も含めます。dimensionが増えるsieve bridgeには別途inverse-problem theoryが必要です。",
    ]),
    ("Slide 23  Computation", [
        "Zhao and ShaoのSection 2.3は、nuisance lawを先に推定してtarget modelへplug-inする順序の先例です。本稿の更新アルゴリズムそのものの先例ではありません。本稿のcriterionは非凸なlatent mixtureを含むため、tensor initialization、複数初期値、制約付き更新を組み合わせます。",
        "具体的にはbridgeを推定してweighted tensorsを作り、anchor tensorを分解してlabelを正規化し、extension kernelsを初期化します。その後、確率単体制約の下で全パラメータを更新し、最大criterionの解を採用します。",
    ]),
    ("Slide 24  Simulation design", [
        "simulationはAllman、Matias and RhodesのSections 3から5に従い、本文式35のr-class、p-feature product mixtureから出発します。Allmanらに従うのはmodel classであり、数値parameterは本研究固有です。p fはlatent class比率、G fはclass別のWの分布、M j fはclass別の項目jの測定核です。",
        "標本サイズは500、Fは2 class、WとY 1、Y 2、Y 3は二値です。Y 1は常時観測とし、Y 2とY 3は自身の値に依存して欠測するMNAR、平均item response rateは80パーセントです。各pairには4カテゴリでfull-rankのMissing IVを置き、100回反復します。",
        "比較は五手法です。第一は欠測前のfull product-mixture likelihoodを使うcomplete-data oracle、第二はignorabilityを置くMAR full-information likelihood、第三は真のlogistic selection equationを用いてW、Y obs、Dの周辺尤度を最大化するcorrect selection likelihood、第四はY 1のmain effectとinteractionを除くmisspecified selection likelihoodです。第五の提案法は、各blockで固定4-cellのfinite-cell saturated inverse bridgeを推定し、weighted pairwise latent-class criterionを当てはめます。",
    ]),
    ("Slide 25  Evaluation targets", [
        "新しい評価対象の一つ目はmeasurement kernel M jです。二値項目ではM j of 1, fは、latent classがfのときに項目jが1となる確率です。これは項目とlatent classの測定関係を表し、全てのitem-by-class cellsにわたる平均biasと平均RMSEを報告します。",
        "二つ目はlatent class proportion p f、つまり母集団におけるclass fの比率です。本文の結果表に合わせてp 2を評価し、p 1は1 minus p 2で決まります。Biasは系統誤差を、RMSEはbiasとreplication variabilityを合わせた有限標本誤差を表します。",
    ]),
    ("Slide 26  Simulation results", [
        "measurement kernelsについて、提案法はMARと誤指定selection likelihoodよりabsolute biasとRMSEが小さく、測定関係をより正確に回復しています。これは、selection equationを直接指定せず、Missing IVからblock lawsを回復しているためと考えられます。",
        "latent class proportionsについても提案法のbiasはほぼゼロですが、RMSEが全比較で常に最小になるわけではありません。正しく指定されたparametric selection likelihoodは効率面で競争的であり、提案法の利点は誤指定回避と引き換えの頑健性にあります。",
    ]),
    ("Slide 27  Conclusion and limits", [
        "第一段階ではMissing IVとcomplete-case completenessにより、parametric selection modelを置かずに各supported-block lawを識別します。第二段階ではanchor itemsとWにより、有限latent-class proportionsと全measurement kernelsを共通labelの下で識別します。推定量もこの順序に沿います。",
        "限界もあります。識別には、既知のclass数r、anchor ordering、supported pairs、Missing IV exclusion、completeness、局所独立性、rank条件、有限latent-class modelが必要です。提案法はselection-link misspecificationには頑健ですが、これらの仮定やlatent-class modelの誤指定まで許すものではありません。また、class proportionのRMSEは正指定selection likelihoodを上回り得ます。漸近正規性は有限次元bridgeに対する結果であり、sieve inferenceには追加のinverse-problem theoryが必要です。",
    ]),
    ("Slide 28  Selected references", [
        "Stage 1はd’HaultfoeuilleとZhao and Shao、complete-case-law上のcompletenessはMiaoら、Stage 2はAllman、Matias and Rhodesに依拠しています。Lewbelはobservational equivalenceとpoint identificationの用語、Little and Rubinは欠測機構、Heckmanはsample-selection biasの基礎として参照しています。",
    ]),
    ("Appendix Slide 29  Steps 1 and 2", [
        "Step 1では命題1によりanchor pairと全extension pairsのfull lawsを回復します。Step 2ではX equals xを固定し、局所独立性からY a、Y b、WをFとX equals xの下で三つの条件付き独立なviewとするanchor tensorを構成します。",
    ]),
    ("Appendix Slide 30  Steps 3 and 4", [
        "Step 3ではKruskal rank conditionからanchor tensorの分解をcommon permutationとcolumn scalingを除いて一意にします。Step 4では確率ベクトルの列和でscalingを固定し、anchor orderingでclass labelsを固定します。",
    ]),
    ("Appendix Slide 31  Steps 5 and 6", [
        "Step 5では補題3.1によりG Khatri–Rao M aがfull column rankであることを示し、extension inversionを一意にします。Step 6では識別されたpとGへBayes則を適用し、WとXを条件とするlatent class distributionを得ます。",
    ]),
    ("Appendix Slide 32  Steps 7 and 8", [
        "Step 7では各extension tensorを線形反転して全measurement kernels M jを識別します。Step 8では識別されたlatent weightsとmeasurement kernelsを局所独立モデルへ代入し、一度も同時観測されていない項目集合のjoint lawを構成します。",
    ]),
    ("Appendix Slide 33  Empirical diagnostics and falsification checks", [
        "候補となるMissing IV、latent shifter、covariates、anchor itemsの割当てについて、relevance、finite complete-case operatorのrank、supported-block positivity、third-mode rank、anchor separationを診断できます。冗長なblocksから得られるmeasurement kernelsの一致はoveridentifying restrictionになります。",
        "ただし、診断に適合するだけでMissing IV exclusion、measurement exclusion、局所独立性が検証されたとはいえません。不適合は候補設計を反証できますが、適合は全仮定の正しさを保証しません。",
    ]),
]


def _slide_by_title(title):
    for heading, paragraphs in SLIDES:
        if heading.split("  ", 1)[1] == title:
            return title, paragraphs
    raise KeyError(title)


_original_slides = SLIDES
SLIDES = [
    _slide_by_title("Title"),
    _slide_by_title("Missing-data mechanisms and standard approaches"),
    _slide_by_title("Usual IV and Missing IV"),
    _slide_by_title("Latent measurement model"),
    _slide_by_title("Related identification results"),
    (
        "Proposed two-stage identification",
        [
            "本稿の方法を先に概観します。第一に、Zhao and Shaoの除外制約とd’Haultfoeuilleのconditional moment equationを各supported blockへ適用し、parametric selection modelを置かずにblock lawを回復します。",
            "第二に、回復したoverlapping block lawsをanchor itemsとlatent shifterで接続し、Allman、Matias and Rhodesのtensor uniquenessとextension inversionを用いて共通label下の潜在分布と測定核へ分解します。本稿の新規性は、この二段階を一つの識別定理として接続する点です。",
        ],
    ),
    _slide_by_title("Model and assumptions"),
    _slide_by_title("A latent shifter is not a Missing IV"),
    ("Assumption 1", [
        "仮定1は観測過程です。Yは全測定項目、Dは各項目を観測したかを示す二値指標です。Y jはD jが1の場合にだけ観測され、X、W、Z、Dは常時観測されます。Fは潜在変数で観測されません。",
        "Oは一個体の観測データを表します。常時観測変数に、実際に観測された項目だけからなるY Dを加えたものです。Zには各blockで必要なMissing IV Z Sが含まれます。",
    ]),
    ("Assumptions 2 and 3", [
        "仮定2では通常共変量をX、潜在分布を動かすlatent shifterをWとします。Q FはWとXを条件とするFの分布です。この分布ではZを周辺化しています。",
        "仮定3は測定過程です。Q jはFとXを条件とする項目Y jの分布、theta jはそのパラメータです。Q Fは潜在classの構成、Q jは各classと測定項目との関係を記述します。",
    ]),
    ("Assumption 4", [
        "仮定4は、Zを周辺化した測定モデルに局所独立性とWのmeasurement exclusionを課します。F、W、Xを条件としたY全体の分布が、Wを含まない各Q jの積になります。全項目の同時分布に対する条件であって、pairwise independenceだけではありません。",
        "ここでF、W、Xの下でYとZが独立とは仮定しません。この強い条件は削除し、ZとYの関連を許します。Missing IVとしてのZのexclusionは、仮定6で欠測指標Rに対して別に課します。この区別によって、測定モデルと四cellのshadow operatorを両立できます。",
    ]),
    _slide_by_title("A simple multivariate extension"),
    ("Assumption 5", [
        "仮定5では、Sを測定項目の部分集合、Y Sをその測定ベクトル、U SをXとW、R SをS内のD jの積と定義します。R Sが1ならblock全体が観測されています。一部だけ観測されたblockは、complete blockとしては用いません。",
        "supported blockとは、対象support上で、Y SとU Sを条件とする観測確率pi Sが正であるblockです。必要なのはanchor pairとextension pairsであり、全項目のglobal complete caseは必要としません。",
    ]),
    ("Two marginal views of the model", [
        "図は、一つの完全データ法則の異なる周辺を二つに分けています。上段ではZとDを周辺化し、Fと測定項目の関係を示します。U SからY Sへの辺はXの効果だけであり、Wは測定核へ入りません。項目間の局所独立性は仮定4の積分解によって課しています。",
        "下段はFとblock外の項目を周辺化した顕在変数のfull-data lawです。Y Sはここでも欠測し得る変数です。Y SとU Sの下でZ SとR Sが独立になる分解を表します。上下の図を一つのDAGとして結合したり、FとZの追加の独立性を読んだりするものではありません。",
    ]),
    _slide_by_title("Overlapping supported blocks replace global complete cases"),
    ("Assumptions 6 and 7", _slide_by_title("Stage-1 assumptions")[1]),
    _slide_by_title("New Proposition 1"),
    _slide_by_title("Proof sketch the observable bridge is unique"),
    _slide_by_title("Finite latent-class identification"),
    _slide_by_title("The remaining decomposition problem"),
    _slide_by_title("Identification strategy"),
    ("Assumption 3.1: Support and model", [
        "仮定3.1の前半は、finite latent-class modelの対象範囲と第一段階との接続です。class数rは既知で2以上、項目空間とXで条件づけたWのsupportは有限です。全ての条件を、Xの共通の確率1の集合上で課します。",
        "仮定2から4が成立し、各class比率とWのsupport上の確率は正です。anchor pairと全extension pairsについて仮定1、5から7を満たし、真の逆観測確率はcomplete-case lawの下で二乗可積分とします。これが命題1を各pairへ適用する条件です。",
    ]),
    ("Assumption 3.1: Rank and labels", [
        "後半はrankとlabelです。M a、M b、Gのcolumn Kruskal rankの和を2rプラス2以上とします。三つのviewに対するこの条件自体はAllmanらのTheorem 1と同じです。",
        "さらに既知のanchor scoreを各classのM aで平均したmu afについて、全てのXで同じ厳密順序を課します。rankが因子を置換を除いて固定し、orderingが共通labelを固定します。GとM aのKhatri–Rao積のfull column rankは、次の補題から導く結論であり、追加仮定ではありません。",
    ]),
    _slide_by_title("Lemma 3.1"),
    _slide_by_title("New Theorem 1"),
    _slide_by_title("Proof map for Theorem 1"),
    _slide_by_title("Bridge-weighted composite estimation"),
    ("Conditions for estimation", [
        "本文の命題4.1の条件も明記します。仮定1から7と3.1に加え、真のthetaが潜在モデルを正しく生成し、全てのpair densityが共通support上で正とします。母集団criterionには真のbridgeと固定された正のblock重みを使います。",
        "IPWによって各pairのfull-data expectationが回復されます。真値とのcriterionの差はKL divergenceの正の加重和なので、最大値が等しければ全pair lawsが一致し、定理1からthetaも一致します。これが推定の母集団での一意性です。",
    ]),
    ("Assumption 5.1: Regularity", [
        "漸近理論では個体の観測データが独立同分布であるとします。betaを潜在パラメータthetaとbridgeパラメータetaの組とします。確率単体制約を除いた自由座標で、両者は有限次元で正しく指定され、パラメータ空間はcompact、真値は内点です。",
        "first-stage population momentの零点は真値だけとします。criterionとmomentは真値近傍で二回連続微分可能で、必要なenvelopeは二乗可積分です。bridgeとcomposite scoreのJacobian、および対応するGodambe行列が非特異であることも仮定します。",
    ]),
    ("Assumption 5.1: Uniform convergence", [
        "一致性には真値の近傍の微分可能性だけでは不十分です。sample criterionは、推定したbridgeをplug-inした状態で、パラメータ空間全体にわたり真のbridgeを用いた母集団criterionへ一様に収束すると仮定します。",
        "また、真値から任意の正の距離だけ離れたパラメータの母集団criterionは、真値での最大値より厳密に小さいとします。この一様収束と分離を用いてargmax theoremを適用することが、一致性の根拠です。",
    ]),
    ("Assumption 5.1: First stage and separation", [
        "仮定5.1の後半では、bridge推定量がregular asymptotically linearであると仮定します。そのinfluence functionは平均ゼロで有限分散を持ちます。",
        "さらに、class比率、supported-pair propensity、anchor tensorのW-mode unfoldingの最小の非零特異値、Kruskal rankを支えるminor、anchor orderingの隣接gapが一様にゼロから離れているとします。unfoldingはWを行、Y aとY bを列とする行列です。識別だけでは推定の正則性や安定した推論までは従わないので、これらを別に明示しています。",
    ]),
    ("Theorem 5.1", [
        "命題4.1と仮定5.1の下で、有限次元bridgeを用いた推定量の一致性と漸近正規性を得ます。Aはcomposite scoreのtheta微分の期待値、Cはeta微分の期待値です。分散BはscoreだけでなくCとfirst-stage influence functionの積を含みます。",
        "したがってsandwich分散にはbridgeを推定した誤差も入ります。個人単位bootstrapでも両段階とlabel alignmentをやり直します。増大するsieveについてはこの定理だけでは足りず、推定速度とinverse-problemの条件が追加で必要です。",
    ]),
    _slide_by_title("Computation"),
    (
        "Simulation targets in the model",
        [
            "ここではシミュレーションで実際に使う生成法則を図示し、評価する母数を書き添えます。Xは固定で、FからWと各Yを条件付き独立に生成し、そのY SからZ SとR Sを生成します。Fは2 class、Wと各Yは二値です。Y 1は常時観測されるため、R SはD jに一致します。",
            "Fの周辺class比率がp f、項目とclassの測定関係がM jfです。Stage 1のmu jは母集団周辺確率、p S,cはpairのcell probabilityです。これらと、Stage 2で分離する六つの測定核成分およびp 2を区別します。標本サイズは500、Monte Carlo反復は100回、平均item response rateは80 percentです。",
            "このDGPは、Zを周辺化した仮定4の測定モデルとMissing IV exclusionを満たします。さらにY Sの下でR SとF、W、Z Sが独立になる具体的な生成例です。この追加の独立性を一般の識別定理の仮定として課しているわけではありません。",
        ],
    ),
    (
        "Simulation design: estimators and evaluation",
        [
            "五手法を比較します。complete-data oracleは欠測前の尤度、MARはignorabilityの下での観測尤度、正指定selection likelihoodは真のlogistic selection equationを用います。誤指定版ではY 1の主効果とinteractionを除きます。提案法は各blockのMissing IV momentから4-cell inverse bridgeを推定し、Stage 1の法則を得た後、weighted pairwise latent-class criterionを用いて潜在構造を推定します。",
            "outcome parameterごとに四つの指標を報告します。Biasは平均的な推定誤差、empirical SDは反復間のばらつき、平均estimated SEは各標本で得た不確実性の推定値、95 percent coverageは真値を含むpointwise Wald区間の割合です。likelihood法ではdelta-method sandwich SE、提案法では各反復で個人単位bootstrapを200回行い、両bridgeを再推定します。",
            "潜在母数についてはbiasとRMSEを報告します。測定核のRMSEは、反復と六つのitem-by-class cellsをまとめた二乗誤差の平均の平方根であり、cell別RMSEの単純平均ではありません。p 2は単独で評価します。提案法のoutcome lawはdecomposition前のnormalized IPW推定値であり、他の四手法では推定モデルから導いた法則を比較しています。",
        ],
    ),
    (
        "Population outcome recovery",
        [
            "横軸を五手法、青い実線とオレンジの破線をmu 2、mu 3とした折れ線グラフです。上段はbiasとempirical SD、下段はmean estimated SEと95 percent coverageを示します。",
            "母集団周辺確率の真値はmu 2 equals 0.475、mu 3 equals 0.4925です。提案Stage-1 bridgeのbiasはそれぞれマイナス0.0099、マイナス0.0096で、empirical SDと平均estimated SEは0.0379対0.0365、0.0345対0.0339でした。95 percent coverageは94 percentと93 percentで、推定SEがreplication variabilityを概ね捉えています。",
            "これに対してMARのcoverageは9 percentと13 percent、誤指定selection likelihoodでは12 percentと11 percentでした。正しく指定したselection likelihoodはbiasがほぼゼロですが、coverageはY 2で86 percentです。本結果は提案法の効率優位を示すものではなく、parametric selection linkを指定せずに、欠測outcomeの周辺分布をStage 1で補正できることを示します。",
        ],
    ),
    (
        "Stage-2 full joint law",
        [
            "続いて、Y 1、Y 2、Y 3のfull joint lawを全体像として示します。横軸は三つのbinary outcomesの8通りの組合せ、縦軸は各cellの確率です。黒い線が母集団の真の分布、そのほかの線は各手法で推定した分布を100反復で平均したものです。",
            "青い提案法の線は、Stage 1で回復したoverlapping block lawsをStage 2で共通の潜在構造に接続し、class比率と全measurement kernelsから構成した分布です。周辺確率だけでなく、全項目の組合せとして真の分布をどのように再現しているかを見ます。",
        ],
    ),
    ("Latent-parameter recovery", _slide_by_title("Simulation results")[1]),
    _slide_by_title("Conclusion and limits"),
    _slide_by_title("Selected references"),
    ("Proof of Theorem 1 (Steps 1--2)", _slide_by_title("Steps 1 and 2")[1]),
    ("Proof of Theorem 1 (Steps 3--4)", _slide_by_title("Steps 3 and 4")[1]),
    ("Proof of Theorem 1 (Steps 5--6)", _slide_by_title("Steps 5 and 6")[1]),
    ("Proof of Theorem 1 (Steps 7--8)", _slide_by_title("Steps 7 and 8")[1]),
    _slide_by_title("Empirical diagnostics and falsification checks"),
    (
        "Simulation calibration",
        [
            "WとYの周辺モデルにはAllman、Matias and Rhodesの有限product mixtureを用い、数値は本研究で設定しています。class比率は0.55と0.45、G fはFを条件としたWの分布です。class labelsはY 1の測定確率の小さい順に固定します。",
            "Y 1は常時観測されます。Y 2とY 3の観測確率は自身、Y 1、およびinteractionに依存するlogistic式で、各項目の平均観測率が70 percentとなるよう切片を調整します。全三項目の平均観測率は80 percentです。Z Sはpairの四つのcellに依存するfull-rank行列Hで生成し、常時観測します。",
            "有限supportを全列挙した母集団で検証すると、各Wの下でcomplete-case operatorのrankは4、三つのviewのKruskal rankの和は6、extension inversionのrankは2です。改訂した仮定4では、FとWの下でY SとZ Sが独立という強い条件は不要です。これらは識別条件の確認であり、推定量の全正則条件を数値的に証明するものではありません。",
        ],
    ),
]
appendix_start = next(
    index for index, (title, _) in enumerate(SLIDES, start=1)
    if title == "Proof of Theorem 1 (Steps 1--2)"
)
SLIDES = [
    (
        f"{'Appendix ' if index >= appendix_start else ''}Slide {index}  {title}",
        paragraphs,
    )
    for index, (title, paragraphs) in enumerate(SLIDES, start=1)
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def set_run_font(run, latin, east_asia, size=None, bold=None):
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = LATIN_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), JAPANESE_FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    title_style = styles["Title"]
    title_style.font.name = LATIN_FONT
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), JAPANESE_FONT)
    title_style.font.size = Pt(20)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.paragraph_format.space_after = Pt(12)
    for border in list(title_style.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)

    for style_name, size in [("Heading 1", 15), ("Heading 2", 12.5)]:
        style = styles[style_name]
        style.font.name = LATIN_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), JAPANESE_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    if "Slide Reference" not in styles:
        ref_style = styles.add_style("Slide Reference", WD_STYLE_TYPE.PARAGRAPH)
        ref_style.font.name = LATIN_FONT
        ref_style._element.rPr.rFonts.set(qn("w:eastAsia"), JAPANESE_FONT)
        ref_style.font.size = Pt(9)
        ref_style.font.color.rgb = RGBColor(80, 80, 80)
        ref_style.paragraph_format.space_after = Pt(3)

    header = section.header.paragraphs[0]
    header.text = "欠測IVと潜在変数モデリングによる非無作為欠測の識別"
    set_run_font(header.runs[0], LATIN_FONT, JAPANESE_FONT, 8.5)
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_page_number(section.footer.paragraphs[0])

    doc.add_paragraph("欠測IVと潜在変数モデリングによる非無作為欠測の識別", style="Title")
    subtitle = doc.add_paragraph("英語Beamer対応 日本語発表原稿")
    for run in subtitle.runs:
        set_run_font(run, LATIN_FONT, JAPANESE_FONT, 12, True)

    doc.add_paragraph(
        f"本原稿は、英語Beamer全{len(SLIDES)}枚に対応する日本語の口頭説明用原稿である。"
        "識別対象、本文の全仮定、定理、推定法と数値結果を説明する。測定モデルはZを周辺化した法則とし、Missing IV exclusionと区別する。"
        "証明は本編では論理の流れ、Appendixでは8ステップを説明する。"
    )

    doc.add_paragraph("発表の構成", style="Heading 1")
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    table.columns[0].width = Inches(1.55)
    table.columns[1].width = Inches(5.1)
    hdr = table.rows[0].cells
    hdr[0].text = "範囲"
    hdr[1].text = "内容"
    for cell in hdr:
        set_cell_shading(cell, "D9E7F5")
        for run in cell.paragraphs[0].runs:
            set_run_font(run, LATIN_FONT, JAPANESE_FONT, 10, True)
    set_repeat_table_header(table.rows[0])
    rows = [
        ("Slides 1–7", "基礎概念、既存研究、二段階の識別戦略"),
        ("Slides 8–18", "仮定1から7、supported-block lawの回復"),
        ("Slides 19–26", "仮定3.1、有限潜在クラスのtensor分解"),
        ("Slides 27–40", "推定、仮定5.1、漸近理論、計算、simulation、結論、参考文献"),
        ("Appendix 41–46", "定理1の8ステップ、empirical diagnostics、数値設定"),
    ]
    for left, right in rows:
        cells = table.add_row().cells
        cells[0].text = left
        cells[1].text = right
        for cell in cells:
            for run in cell.paragraphs[0].runs:
                set_run_font(run, LATIN_FONT, JAPANESE_FONT, 10.5)

    doc.add_page_break()
    for index, (heading, paragraphs) in enumerate(SLIDES):
        doc.add_paragraph(heading, style="Heading 1")
        for text in paragraphs:
            paragraph = doc.add_paragraph(text)
            paragraph.paragraph_format.keep_together = True
        if index < len(SLIDES) - 1:
            spacer = doc.add_paragraph()
            spacer.paragraph_format.space_after = Pt(1)

    doc.core_properties.title = "欠測IVと潜在変数モデリングによる非無作為欠測の識別 日本語発表原稿"
    doc.core_properties.author = "Masahiro Honda"
    doc.core_properties.subject = "English Beamer presentation script in Japanese"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
