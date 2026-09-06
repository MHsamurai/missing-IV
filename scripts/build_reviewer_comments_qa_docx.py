from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "manuscript" / "vector_missing_iv_reviewer_comments_qa_ja.docx"
LATIN_FONT = "Arial"
JAPANESE_FONT = "BIZ UDPGothic"


MAJOR_COMMENTS = [
    {
        "title": "1. 貢献の境界を、二段階の接続としてさらに固定する",
        "priority": "最重要",
        "location": "第1節、とくに既存研究の整理と最終2段落",
        "comment": (
            "本稿の新規性は、latent structureとMNARを同時に扱うこと自体ではない。"
            "Lee and Tang、Holman and Glas、Kuha et al.、Jung et al.などは、"
            "欠測機構とlatent modelをjointに規定している。また、第一段階のshadow-variable identificationと"
            "第二段階のthree-view tensor decompositionも、それぞれ既存の原理である。"
            "本稿固有の貢献は、parametric selection modelを規定せずにsupported-block lawsを回復し、"
            "overlap、anchor items、latent shifterを用いて有限潜在クラス比率と測定核を共通labelの下で"
            "識別する接続にある。現稿は概ねこの位置づけになっているが、応用例からlatent-MNAR文献へ移る箇所では、"
            "読者が『latent-MNARそのものが未研究』と誤読する余地が残る。"
        ),
        "request": (
            "abstract、Introduction、Conclusionで、同じ一文により貢献を固定する。"
            "『初めて』『既存研究では識別されていない』という広い主張は避け、"
            "selection-model-free block-law recoveryとlatent decompositionの統合に限定する。"
        ),
        "response": (
            "ご指摘に従い、新規性を二段階の識別原理の接続として明確化しました。"
            "既存のlatent-MNAR joint modellingを明示的に認めた上で、"
            "shadow IVとcomplete-case completenessによるsupported-block lawの回復、"
            "ならびにanchor itemsとlatent shifterによる共通label下の分解が本稿の差であると統一しました。"
        ),
    },
    {
        "title": "2. 命題2.1の関数空間、support、bridge存在条件を明示する",
        "priority": "最重要",
        "location": "仮定5--7、命題2.1、式(11)--(17)",
        "comment": (
            "complete-case completenessからinverse bridgeの一意性を導く論理は明快である。"
            "ただし、査読では、どのsupport上のどの作用素が単射なのか、条件がuごとなのかjoint law上なのか、"
            "正値解qの候補集合は何か、真のpiの逆数が二乗可積分であること以外にbridgeの存在を仮定しているのか、"
            "という点が問われる。さらに、P(R_S=1|Z_S,U_S)>0で除す段階が、"
            "式(11)のpositivityからどの意味で従うかをsupportの言葉で説明した方がよい。"
        ),
        "request": (
            "conditional expectation operatorを明示的に定義し、そのdomain、codomain、a.s.の基準測度を示す。"
            "命題の直前にcandidate bridge classを定義し、positivityがcomplete-case lawとfull lawの"
            "零集合を一致させる範囲を記す。連続変数を許す場合はregular conditional distributionの存在も整理する。"
        ),
        "response": (
            "命題2.1の候補集合を、正値かつinverseがcomplete-case lawの下でL2に属する関数として定義し、"
            "complete-case conditional operatorの単射性をその集合の差に適用する形へ書き換えます。"
            "また、positivityにより必要な条件付き確率が正となり、complete-case law上の等式が"
            "full-law supportへ移ることを証明中に明示します。"
        ),
    },
    {
        "title": "3. 定理3.1の識別概念と正規化を、定理内で完結させる",
        "priority": "最重要",
        "location": "第3節、式(19)--(27)、仮定3.1、定理3.1",
        "comment": (
            "Lewbelの意味でのobservational equivalence、商空間上の識別、anchor ordering後のpoint identificationを"
            "区別した点は適切である。ただし重要な限定であるknown class number r、共通full-probability set X0、"
            "strict anchor ordering、positive class probabilitiesが仮定3.1に集約されているため、"
            "定理だけを読むと識別範囲を広く受け取られやすい。"
        ),
        "request": (
            "定理文に『既知r』『x in X0ごと』『anchor orderingで正規化したTheta_x上』を残し、"
            "orderingを課さない場合の結論を直後のremarkとして示す。"
            "global point identificationとAllman et al.のgeneric identificationを本文と付録で一貫して区別する。"
        ),
        "response": (
            "識別対象を正規化されたparameter space上のpoint identificationとして定理文に明記し、"
            "orderingがない場合はclass permutationによる商空間上の識別であることを別記します。"
            "本結果がrealized Kruskal rankに基づくglobal resultであり、generic identificationではないことも維持します。"
        ),
    },
    {
        "title": "4. extension kernelsの回復条件をblockごとに検証できる形で示す",
        "priority": "重要",
        "location": "仮定3.1、定理3.1のStep 5--7、式(26)--(27)",
        "comment": (
            "anchor tensorの一意分解からG tensor M_aのfull column rankを導き、各extension tensorを反転する流れは"
            "本稿の中核である。したがって、全てのnonanchor itemについてextension pair {j,a}がsupportedであること、"
            "同じmeasurement kernelとclass labelsを共有すること、rank deficiencyやnear-singularityがある場合の"
            "結論が明確でなければならない。現稿では条件は書かれているが、識別定理の直感との対応がやや埋もれている。"
        ),
        "request": (
            "仮定3.1を、anchor decompositionに必要な条件とextension propagationに必要な条件へ視覚的に分ける。"
            "式(27)の右逆が一意なM_jを与える理由と、零確率cellや小さい特異値が識別・推定へ与える影響をremarkで補う。"
        ),
        "response": (
            "anchor pairに関するKruskal条件と、各extension pairに関するsupportednessおよびlinear inversion条件を"
            "分けて提示します。識別にはfull column rank、数値安定性には最小特異値が必要であることを区別し、"
            "後者をalgorithmのdiagnosticへ接続します。"
        ),
    },
    {
        "title": "5. 有限次元bridgeの理論とsaturated/sieve実装の距離を埋める",
        "priority": "最重要",
        "location": "第4--6節、仮定5.1、定理5.1、simulationのproposed method",
        "comment": (
            "定理5.1は有限次元で正しく指定されたbridgeとlatent modelを対象とする一方、本文の提案法は"
            "sieve、cross-fitting、regularizationにも言及し、simulationではcell-saturated inverse bridgeを用いる。"
            "有限カテゴリで固定cell数ならsaturated bridgeは有限次元であるが、『sieve』と呼ぶとdimension growthと"
            "ill-posed inverse problemに対するrate theoryまで示したように読まれる。"
        ),
        "request": (
            "主要定理の適用範囲をfinite-dimensional bridgeに限定し、simulationの手法は"
            "finite-cell saturated bridgeと呼ぶ。sieve extensionは明確に今後の課題または補足的議論とし、"
            "cross-fittingを理論に使うならsample splitting、fold数、influence functionへの反映を示す。"
        ),
        "response": (
            "漸近正規性はfinite-dimensional bridgeに対する結果であることを本文・図表・結論で統一します。"
            "simulationの実装は固定4-cellのsaturated bridgeと記し、growing sieveに対する通常のroot-n theoryは"
            "主張しません。sieve推論にはsource condition等が必要である点を限界として残します。"
        ),
    },
    {
        "title": "6. composite criterionの推論で第一段階誤差を完全に扱う",
        "priority": "重要",
        "location": "式(30)--(34)、命題4.1、仮定5.1、定理5.1",
        "comment": (
            "母集団criterionの一意性をKL divergenceと定理3.1へ接続する議論は妥当である。"
            "一方、実際の標準誤差には、複数blockで同一個体が繰り返し寄与する相関、bridge推定誤差、"
            "label alignment、境界に近いclass probabilityが影響する。Godambe行列を述べるだけでは実装の再現性が不足する。"
        ),
        "request": (
            "stacked estimating equationsの具体形、A・B・Cの標本推定量、block間相関を保持する個体単位cluster、"
            "bootstrap時の全段階再推定をalgorithmまたは付録に記載する。"
            "anchor orderingが局所的に不変となるseparation条件の役割も説明する。"
        ),
        "response": (
            "first-stage momentsとsecond-stage composite scoreを個体単位でstackし、"
            "同一個体内のblock contributionsをまとめたsandwich varianceを提示します。"
            "bootstrapではbridge、latent fitting、label alignmentを全て再実行する手順を明記します。"
        ),
    },
    {
        "title": "7. simulationを識別条件の強弱に結びつける",
        "priority": "最重要",
        "location": "第7節、図4、表1、補題7.1",
        "comment": (
            "五手法の比較は、selection equationの誤指定に対する提案法のbias低減を示しており有益である。"
            "しかし現在はn=500、r=2、binary features、平均item観測率80%、強い4-category Missing IVという一設定である。"
            "本稿の理論が強調するcompleteness、positivity、Kruskal rank、anchor separationが弱くなった場合の"
            "有限標本挙動は確認されていない。また100 replicationsではtail behaviorやcoverageの評価には限界がある。"
        ),
        "request": (
            "少なくともMissing IVの強さ、最低pair propensity、class separationまたはanchor gap、nのいずれかを変えた"
            "感度分析を追加する。推論を主張するならstandard errorと95% coverageを報告する。"
            "M_jの集約biasが符号相殺を起こし得るため、mean absolute biasまたはcomponent-wise結果も補足する。"
        ),
        "response": (
            "主結果ではM_jとp_2のbias・RMSEを維持し、補足でIV strength、overlap、anchor separationの感度分析を追加します。"
            "component-wise biasを併記して集約値の符号相殺を確認し、漸近推論を評価する版ではcoverageも報告します。"
        ),
    },
    {
        "title": "8. 『頑健性』の対象をselection linkの誤指定に限定する",
        "priority": "重要",
        "location": "abstract、第7節の結果解釈、Conclusion",
        "comment": (
            "提案法はparametric selection likelihoodのlink misspecificationを避ける点で頑健である。"
            "しかし、Missing IV exclusion、positivity、complete-case completeness、local independence、known r、"
            "finite latent-class model、anchor orderingの誤指定には頑健ではない。"
            "『nonparametric』『robust』という語だけが前面に出ると、latent modelまで無指定と誤読される。"
        ),
        "request": (
            "『parametric selection modelを直接規定しない』『selection-link misspecificationに対するrobustness』と限定する。"
            "正指定parametric likelihoodより一般に効率的ではないことを、結果と結論の双方で維持する。"
        ),
        "response": (
            "robustnessの対象をselection equationのparametric misspecificationに限定し、"
            "latent measurement modelとidentification assumptionsは引き続き必要であることを明記します。"
            "効率性とのtrade-offもsimulation結果に沿って記述します。"
        ),
    },
    {
        "title": "9. 補題7.1を『反証可能な診断』として位置づける",
        "priority": "重要",
        "location": "補題7.1とその証明概略",
        "comment": (
            "候補となるMissing IV、latent shifter、anchorの割当てについて、relevance、operator rank、positivity、"
            "third-mode rank、anchor gapをデータから診断する発想は応用上重要である。"
            "ただしexclusion、measurement exclusion、local independenceは一般に観測データだけでは検証できない。"
            "この補題を『割当てが正しいことを検証できる』と表現すると過大である。"
        ),
        "request": (
            "necessary diagnosticsまたはfalsification checksと呼び、成立は十分条件の真実性を保証しないことを強調する。"
            "冗長blockによるoveridentifying restrictionsと、識別仮定そのものを区別する。"
        ),
        "response": (
            "補題の役割を候補設計の反証可能性に限定します。観測可能なrank・overlap・anchor gapと、"
            "観測不能なexclusion・local independenceを明確に分けます。"
        ),
    },
    {
        "title": "10. SEM・IRTへの含意と有限潜在クラス定理の範囲を分ける",
        "priority": "重要",
        "location": "第1節、定理3.1の脚注、Conclusion、発表スライド",
        "comment": (
            "SEMやIRTは動機として適切であるが、主定理はfinite latent classを対象とする。"
            "連続因子では位置・尺度・符号の正規化だけでなく、非線形再パラメータ化を除く追加構造が必要である。"
            "読者が本定理をlinear factor analysisやRasch modelへ直接適用できると理解しないよう、"
            "現時点の対象と将来拡張を明確に分ける必要がある。"
        ),
        "request": (
            "IntroductionではSEM・IRTをmotivationとし、formal contributionはfinite latent-class measurement modelと明記する。"
            "continuous factor、linear factor model、Rasch modelへの適用は別定理を要することをConclusionで再確認する。"
        ),
        "response": (
            "SEM・IRTは応用上の動機として残し、現稿の識別定理が有限潜在クラスに限定されることを"
            "abstract、theorem、conclusionで統一します。連続因子への拡張は追加のcalibrationまたはoperator injectivityを"
            "必要とする別課題として扱います。"
        ),
    },
]


MINOR_COMMENTS = [
    ("用語", "最初にMissing IV / shadow variableと併記した後はMissing IVへ統一し、通常のIVと明確に区別する。"),
    ("記号", "O、Y_D、R_S、U_S、Z_S、p_f、lambda_f、G、M_jは初出で文章による定義も付ける。"),
    ("条件づけ", "X=xを固定する議論と、Xを含む確率法則の議論を混在させない。a.s.の基準となる分布も明記する。"),
    ("式番号", "displayed equationsには番号を付け、本文では『上式』ではなく番号で参照する。"),
    ("定理名", "本文とBeamerで命題・定理のheadingを一致させる。BeamerのNew!は発表用に限定する。"),
    ("simulation表", "p_fの結果がp_2を指すことを表頭または注で明示し、M_jの集約方法を再現可能に定義する。"),
    ("図", "DAGは条件付き独立構造、Figure 3は識別手順を示す図として役割を分け、caption単独で解釈できるようにする。"),
    ("引用", "AllmanのTheorem 1/Corollary 2とTheorem 4、Zhao--ShaoのSections 2.3と4を混同しない。"),
    ("Miao et al.", "working-paper版を引用する場合は版を明示し、主要貢献の根拠をpublished literatureでも補えるか確認する。"),
    ("再現性", "乱数seed、software、optimizer、収束判定、初期値生成、ridge値、constraint処理をsimulation appendixに記載する。"),
]


QUESTIONS = [
    (
        "Q1. 本稿の最も新しい点は何ですか。",
        "A. Missing IVによるMNAR law recoveryとfinite latent-class decompositionを、supported blocksを介して一つの識別手順へ接続した点です。",
        "第一段階ではparametric selection modelを置かずに各supported-block lawを回復します。第二段階ではanchor itemsとlatent shifter Wをthree viewsとして用い、class proportionsとmeasurement kernelsを共通labelで識別します。個々の原理ではなく、この接続が貢献です。",
    ),
    (
        "Q2. latent-MNAR modelは既に多数あるのではないですか。",
        "A. あります。本稿はlatent-MNAR joint modelling自体を新規とは主張しません。",
        "Lee and Tang、Holman and Glas、Kuha et al.、Jung et al.などはselectionまたはresponse-propensity modelとlatent modelを同時に規定します。本稿との差は、selection equationをparametrically指定せず、Missing IVとcompletenessからblock lawsを先に回復する点です。",
    ),
    (
        "Q3. WをMissing IVとして使えないのですか。",
        "A. 一般には使えません。Wはlatent shifterであり、Fの分布を動かすためです。",
        "Fまで条件づければDとWが独立でも、Fを周辺化するとP(F|Y,W,X)が残ります。WがFを予測すればD independent W given Y,Xは通常成立しません。Z_Sはblock-law recovery、Wはlatent decompositionのthird modeという別の役割です。",
    ),
    (
        "Q4. 通常のIVとMissing IVは何が違いますか。",
        "A. 通常のIVは内生変数を動かしoutcome equationから除外されます。Missing IVは欠測し得るYを予測し、Yと共変量を条件づけるとmissingness indicatorから除外されます。",
        "本稿の条件はR_S independent Z_S given Y_S,U_Sです。Z_Sにはrelevanceとexclusionが必要で、通常のcausal IVと同じ変数役割ではありません。",
    ),
    (
        "Q5. なぜglobal complete caseが不要なのですか。",
        "A. 定理に必要なのはanchor pairと各extension pairがそれぞれ正の確率で観測されることだからです。",
        "全項目が一度に観測されなくても、overlapping supported pairsを回復し、共通anchorを介してkernelsとlabelsを伝播できます。global complete caseの確率が0でも、必要なpair probabilitiesが正なら識別可能です。",
    ),
    (
        "Q6. complete-case completenessは何を保証しますか。",
        "A. observable bridge momentの正値解を一意にします。",
        "別のinverse bridgeと真のinverse propensityの差は、complete casesでZ_S,U_Sを条件づけた期待値が0になります。conditional operatorがcompleteなら差は0であり、positivityによりfull-law supportへ拡張できます。",
    ),
    (
        "Q7. completenessはデータから検証できますか。",
        "A. 一般には完全には検証できませんが、有限カテゴリでは対応するconditional probability matrixのrankを診断できます。",
        "rank failureは候補設計を反証します。ただし有限標本でfull rankに見えることがpopulation completenessを保証するわけではなく、weak singular valuesは推定を不安定にします。",
    ),
    (
        "Q8. proposition 2.1ではbridgeの存在も証明していますか。",
        "A. 真のselection probability pi_Sが存在し、そのinverseが指定したL2 classに属することを仮定しています。命題が新たに示す中心は、そのclass内での一意性です。",
        "存在と一意性を区別し、candidate classとintegrabilityを命題の前後で明示する必要があります。",
    ),
    (
        "Q9. Allman et al.と同じ条件はどこですか。",
        "A. anchor tensorを構成する三つのviewのKruskal rank条件です。",
        "本稿ではY_a、Y_b、WがFとX=xの下で条件付き独立なthree viewsとなり、k_a+k_b+k_G >= 2r+2を課します。違いは、tensor自体をMNAR観測データから第一段階で回復する点と、Wをthird modeにする点です。",
    ),
    (
        "Q10. class label swappingはどのように処理しますか。",
        "A. anchor itemの既知scoreに対するclass-specific meanをstrict orderingし、共通labelを固定します。",
        "orderingがなければ識別はclass permutationを除く商空間上です。orderingを課した正規化parameter space上ではpoint identificationとなります。",
    ),
    (
        "Q11. class数rは識別されますか。",
        "A. 現在の定理ではrは既知です。",
        "class-number selectionは本定理の範囲外です。実装ではinformation criterionやrank diagnosticsを併用できますが、そのselection uncertaintyを含む理論は別途必要です。",
    ),
    (
        "Q12. なぜWをthird modeにする必要がありますか。",
        "A. 二つのanchor itemsだけではmatrix factorizationの一意性が不足するため、class compositionを動かす第三の観測viewが必要です。",
        "G(w,f;x)=P(W=w|F=f,X=x)が十分なcolumn variationを持つと、anchor pairと合わせてthree-way tensorのKruskal uniquenessを利用できます。",
    ),
    (
        "Q13. 一度も同時観測されない項目のjoint lawまで本当に識別できますか。",
        "A. finite latent-class local-independence modelが正しければ識別できます。",
        "各M_jとlambda_f(w,x)を回復した後、sum_f lambda_f product_j M_jに代入するためです。これは観測デザインだけによるnonparametric recoveryではなく、識別されたlatent measurement modelによる外挿です。",
    ),
    (
        "Q14. 提案推定量はfull likelihoodですか。",
        "A. いいえ。bridge-weighted composite M-estimatorです。",
        "各supported pairのfull-law expectationをIPWで再現し、共通latent parametersが全pair lawsへ同時に適合するcriterionを最大化します。block contributionsの依存とfirst-stage errorを含むsandwich varianceが必要です。",
    ),
    (
        "Q15. 提案法はdoubly robustですか。",
        "A. 現在の推定量についてdoubly robustとは主張していません。",
        "Missing IV exclusion、positivity、completeness、正しいfinite latent-class measurement modelが必要です。強みはparametric selection linkを直接指定しないことであり、outcome modelまたはselection modelの一方が正しければよいという二重頑健性ではありません。",
    ),
    (
        "Q16. 正しく指定されたselection likelihoodより優れていますか。",
        "A. biasへの頑健性では利点がありますが、効率性で常に上回るとは主張しません。",
        "simulationではmeasurement-kernel biasは正指定likelihoodに近く、誤指定likelihoodより小さい一方、class proportionのRMSEにはbridge推定とcomposite fittingによる効率損失が残ります。",
    ),
    (
        "Q17. simulationでM_jとp_2を評価する理由は何ですか。",
        "A. M_jは各項目がlatent classをどう測定するか、p_2は母集団のlatent class構成比を表すためです。",
        "full-data marginal meanだけでは、この二つの誤差を区別できません。二値2-class modelではp_1=1-p_2なのでp_2を報告すればclass proportion全体を代表できます。",
    ),
    (
        "Q18. 提案法のsimulation結果を一言で言うと何ですか。",
        "A. selection equationを誤指定した手法やMARよりmeasurement-kernel biasを抑え、正指定selection likelihoodに近い回復を得ています。",
        "ただしoracleよりRMSEは大きく、p_2のRMSEは正指定selection likelihoodより高いです。結果はrobustness-efficiency trade-offとして解釈します。",
    ),
    (
        "Q19. Missing IVが弱い場合はどうなりますか。",
        "A. population identificationが保たれてもinverse problemが不安定になり、bridge varianceとregularization biasが増えます。",
        "finite categorical caseではconditional probability matrixの小さいsingular valueとして現れます。strength、effective sample size、weight distribution、moment residualを必ず報告する必要があります。",
    ),
    (
        "Q20. 変数をMissing IV、latent shifter、anchorへ割り当てた妥当性は検証できますか。",
        "A. 一部は反証可能ですが、全ては検証できません。",
        "relevance、finite-operator rank、positivity、third-mode rank、anchor gap、redundant blocksの整合性は診断できます。一方、shadow exclusion、measurement exclusion、local independenceは、それらの診断が通っても保証されません。",
    ),
    (
        "Q21. continuous latent factorへ拡張できますか。",
        "A. 現在の定理をそのまま拡張することはできません。",
        "平均0、分散1、符号固定だけでは非線形再パラメータ化が残ります。measurement familyの制約、calibration、またはoperator-level injectivityを用いる別の識別定理が必要です。",
    ),
    (
        "Q22. SEMやIRTへの適用をどこまで主張できますか。",
        "A. 現稿ではmotivationと有限latent-class analogueまでです。",
        "linear factor analysisやRasch modelを直接扱うには、それぞれのmeasurement equationとidentification constraintsを組み込んだcorollaryまたは別定理が必要です。",
    ),
    (
        "Q23. DAGの矢印は因果効果を意味しますか。",
        "A. 現稿では主としてjoint-law factorizationとconditional independencesを可視化しています。",
        "因果解釈を与えるには時間順序、介入、exogeneityなど追加の実質仮定が必要です。質疑では『識別に用いる統計構造の図』と説明するのが安全です。",
    ),
    (
        "Q24. 実証応用がなくても統計論文として十分ですか。",
        "A. 識別論文として成立し得ますが、実装可能性を示す補助例があると説得力は上がります。",
        "最低限、simulationで弱いIV、低overlap、近接classを扱い、diagnosticsとfailure modesを示すことが重要です。実データ応用を加える場合はMissing IVとanchorの実質的根拠を明示する必要があります。",
    ),
]


QA_CHECKLIST = [
    ("貢献", "二段階の接続として概ね明確", "abstract・Introduction・Conclusionで同一表現に固定"),
    ("引用監査", "主要文献は原文箇所と照合済み", "Miao working-paper版と最新書誌を最終確認"),
    ("命題2.1", "一意性の証明は通っている", "operator domain、support、bridge classを明記"),
    ("定理3.1", "8-step proofは論理的に接続", "known r、ordering、X0、global/genericの区別を再点検"),
    ("推定理論", "有限次元bridgeの定理あり", "saturated bridgeとgrowing sieveの呼称・範囲を分離"),
    ("分散推定", "sandwichとbootstrapの方針あり", "標本版stacked equationsを付録化"),
    ("simulation", "五手法、M_j・p_2のbias/RMSEあり", "弱識別・overlap感度、必要ならcoverageを追加"),
    ("再現性", "DGPと主要 tuning は本文に記載", "seed、software、optimizer、コード実行手順を固定"),
    ("図表", "三種類の図の役割は分離", "PDFで矢印、凡例、はみ出し、captionを再確認"),
    ("本文・Beamer・Word", "主要用語と結果は同期", "最終TeX修正後に全成果物を再同期"),
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = LATIN_FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), JAPANESE_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)
    if italic is not None:
        run.italic = italic


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])


def add_label_paragraph(doc, label, body):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_together = False
    run = paragraph.add_run(label)
    set_run_font(run, bold=True, color=(28, 67, 121))
    run = paragraph.add_run(body)
    set_run_font(run)
    return paragraph


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_priority_badge(paragraph, priority):
    run = paragraph.add_run(f"[{priority}] ")
    color = (192, 0, 0) if priority == "最重要" else (28, 67, 121)
    set_run_font(run, bold=True, color=color)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = LATIN_FONT
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), JAPANESE_FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.13

    title = styles["Title"]
    title.font.name = LATIN_FONT
    title._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), JAPANESE_FONT)
    title.font.size = Pt(20)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(8)

    for name, size, before in (("Heading 1", 15, 15), ("Heading 2", 12.5, 11), ("Heading 3", 11, 8)):
        style = styles[name]
        style.font.name = LATIN_FONT
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), JAPANESE_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    if "Metadata" not in styles:
        metadata = styles.add_style("Metadata", WD_STYLE_TYPE.PARAGRAPH)
        metadata.font.name = LATIN_FONT
        metadata._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), JAPANESE_FONT)
        metadata.font.size = Pt(9)
        metadata.font.color.rgb = RGBColor(89, 89, 89)
        metadata.paragraph_format.space_after = Pt(3)


def add_summary_table(doc):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (Inches(1.25), Inches(1.3), Inches(4.05))
    headers = ("区分", "評価", "要点")
    for i, (cell, header) in enumerate(zip(table.rows[0].cells, headers)):
        cell.width = widths[i]
        cell.text = header
        set_cell_shading(cell, "D9E7F5")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=9.5, bold=True)
    set_repeat_table_header(table.rows[0])

    rows = [
        ("理論的貢献", "強い", "supported-block recoveryとfinite latent decompositionの接続が明確。"),
        ("証明", "要補強", "主要論理は通るが、operator domain、support、bridge classを形式化したい。"),
        ("推定理論", "限定的", "finite-dimensional bridgeには対応。growing sieve theoryは未提示。"),
        ("数値検証", "有望", "誤指定selection likelihoodに対するbias低減を確認。設計感度は要追加。"),
        ("現段階の判定", "Major revision", "核となるアイデアは維持し、定理の範囲と数値的根拠を締める。"),
    ]
    for row_values in rows:
        cells = table.add_row().cells
        for i, (cell, value) in enumerate(zip(cells, row_values)):
            cell.width = widths[i]
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell)
            for run in cell.paragraphs[0].runs:
                set_run_font(run, size=9.2, bold=(i == 1))
    doc.add_paragraph()


def add_checklist_table(doc):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (Inches(1.25), Inches(2.3), Inches(3.05))
    headers = ("項目", "現状", "最終対応")
    for i, (cell, header) in enumerate(zip(table.rows[0].cells, headers)):
        cell.width = widths[i]
        cell.text = header
        set_cell_shading(cell, "D9E7F5")
        set_cell_margins(cell)
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=9.3, bold=True)
    set_repeat_table_header(table.rows[0])

    for item, status, action in QA_CHECKLIST:
        cells = table.add_row().cells
        for i, (cell, value) in enumerate(zip(cells, (item, status, action))):
            cell.width = widths[i]
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell)
            for run in cell.paragraphs[0].runs:
                set_run_font(run, size=9)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    configure_styles(doc)

    header = section.header.paragraphs[0]
    header.text = "査読コメントおよび想定Q&A"
    set_run_font(header.runs[0], size=8.5, color=(89, 89, 89))
    add_page_number(section.footer.paragraphs[0])

    doc.add_paragraph("欠測IVと潜在変数モデリングによる非無作為欠測の識別", style="Title")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = subtitle.add_run("査読者としてのコメントおよび想定Q&A")
    set_run_font(run, size=13, bold=True, color=(28, 67, 121))
    meta = doc.add_paragraph(style="Metadata")
    meta.add_run(f"レビュー日: {date.today().isoformat()}  |  対象: TeX本文・英語Beamer・日本語発表原稿")

    doc.add_paragraph("1. 総合評価", style="Heading 1")
    doc.add_paragraph(
        "本稿は、多次元MNARの識別を、(i) Missing IVとblockwise complete-case completenessによる"
        "supported-block lawの回復、(ii) anchor itemsとlatent shifterを用いるfinite latent-class decomposition、"
        "という二段階へ分けている。この構成は明確であり、global complete caseを要求せず、"
        "selection equationのparametric specificationに依存しない点に理論的・応用的な価値がある。"
    )
    doc.add_paragraph(
        "中心命題と主定理の論理は相互に接続されている。一方、掲載可能性を高めるには、"
        "第一段階の作用素とsupportの形式化、第二段階の識別概念と正規化、有限次元bridgeの理論範囲、"
        "およびsimulationによる弱識別・overlapへの感度の四点を補強する必要がある。"
        "現段階の内部判定はMajor revisionとするが、核となる識別戦略の変更を求めるものではない。"
    )
    add_summary_table(doc)

    doc.add_paragraph("2. Major Comments", style="Heading 1")
    for item in MAJOR_COMMENTS:
        heading = doc.add_paragraph(style="Heading 2")
        add_priority_badge(heading, item["priority"])
        run = heading.add_run(item["title"])
        set_run_font(run, bold=True)
        add_label_paragraph(doc, "該当箇所: ", item["location"])
        add_label_paragraph(doc, "査読コメント: ", item["comment"])
        add_label_paragraph(doc, "要求する修正: ", item["request"])
        add_label_paragraph(doc, "回答案: ", item["response"])

    doc.add_paragraph("3. Minor Comments", style="Heading 1")
    for label, text in MINOR_COMMENTS:
        paragraph = doc.add_paragraph(style="List Number")
        run = paragraph.add_run(f"{label}: ")
        set_run_font(run, bold=True, color=(28, 67, 121))
        run = paragraph.add_run(text)
        set_run_font(run)

    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("4. 想定Q&A", style="Heading 1")
    doc.add_paragraph(
        "以下は、教授への説明、学会発表、査読回答で問われる可能性が高い論点である。"
        "最初の一文を短答、その後を補足説明として使う。"
    )
    for question, short_answer, detail in QUESTIONS:
        doc.add_paragraph(question, style="Heading 2")
        add_label_paragraph(doc, "短答: ", short_answer)
        add_label_paragraph(doc, "補足: ", detail)

    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("5. 提出前QAチェック", style="Heading 1")
    doc.add_paragraph(
        "最終TeXをfixした後、本文、Beamer、Word原稿を次の順序で再確認する。"
        "理論上の主張とsimulation上の実装を同じ語で呼んでいるかを重点的に見る。"
    )
    add_checklist_table(doc)

    doc.add_paragraph("優先順位", style="Heading 2")
    add_bullet(doc, "P0: 命題2.1の作用素・support・bridge class、定理3.1の識別空間を形式化する。")
    add_bullet(doc, "P0: finite-dimensional bridge theoremとsaturated/sieveの呼称を整理する。")
    add_bullet(doc, "P1: weak Missing IV、low overlap、small anchor gapの感度分析を追加する。")
    add_bullet(doc, "P1: stacked sandwichまたはfull bootstrapの実装記述を付録へ置く。")
    add_bullet(doc, "P2: 実証応用、continuous factor、class-number selectionは別稿または将来課題として扱う。")

    doc.add_paragraph("査読回答の原則", style="Heading 2")
    doc.add_paragraph(
        "回答では、指摘を受け入れたか否かを最初に明示し、変更した節・式・定理を特定する。"
        "本稿が主張していない範囲まで防御せず、finite latent class、known r、selection-link robustnessという"
        "現在の射程を維持する。識別条件の強さに関する指摘には、仮定を弱いと主張するのではなく、"
        "各条件が証明のどこで必要か、どの部分がデータから反証可能かを答える。"
    )

    doc.core_properties.title = "査読コメントおよび想定Q&A"
    doc.core_properties.subject = "欠測IVと潜在変数モデリングによる非無作為欠測の識別"
    doc.core_properties.author = "OpenAI Codex"
    doc.core_properties.keywords = "MNAR, Missing IV, latent class, reviewer comments, Q&A"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
